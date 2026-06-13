"""Resume parsing: extract plain text from uploaded PDF/DOCX files.

PRD §7.1 specifies pdfplumber / python-docx. We add a pypdf fallback for PDFs
that pdfplumber struggles with. Scanned/image-only PDFs yield little or no text;
the caller is expected to detect that (see `looks_empty`) and prompt the user to
re-upload a text-based file (OCR is a future enhancement per the §9 risk table).
"""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)


class UnsupportedFileError(Exception):
    """Raised when the uploaded file type is not a resume we can parse."""


def _extract_pdf(path: Path) -> str:
    text_parts: list[str] = []

    # Primary: pdfplumber (good layout handling). We catch BaseException because
    # the underlying native deps (pdfminer/cryptography) can raise non-Exception
    # errors (e.g. pyo3 PanicException) on some platforms — fall through to pypdf.
    try:
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    text_parts.append(page_text)
    except BaseException as exc:  # noqa: BLE001
        logger.warning("pdfplumber failed for %s: %s", path.name, exc)

    if _join(text_parts).strip():
        return _join(text_parts)

    # Fallback: pypdf.
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text_parts = [(page.extract_text() or "") for page in reader.pages]
    except BaseException as exc:  # noqa: BLE001
        logger.warning("pypdf fallback failed for %s: %s", path.name, exc)

    return _join(text_parts)


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]

    # Pull text out of tables too (resumes often use them for layout).
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return _join(parts)


def _join(parts: list[str]) -> str:
    return "\n".join(p.strip() for p in parts if p and p.strip())


def extract_text(path: Path) -> str:
    """Extract plain text from a resume file. Raises UnsupportedFileError."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix in (".docx",):
        return _extract_docx(path)
    if suffix == ".doc":
        raise UnsupportedFileError(
            "Legacy .doc files aren't supported — please upload a PDF or .docx."
        )
    raise UnsupportedFileError(
        "Unsupported file type. Please upload a PDF or DOCX resume."
    )


def looks_empty(text: str, min_chars: int = 80) -> bool:
    """Heuristic: a scanned/image-only PDF yields little extractable text."""
    return len((text or "").strip()) < min_chars
